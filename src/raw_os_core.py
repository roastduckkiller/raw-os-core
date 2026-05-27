#!/usr/bin/env python3
"""Raw OS v0 core.

This module is intentionally self-contained and stdlib-only for the first
implementation pass. It does not touch any OpenClaw runtime directory unless a
caller explicitly points config.agent.workspaceRoot there.
"""

from __future__ import annotations

import argparse
import importlib.util
import hashlib
import json
import mimetypes
import os
import re
import shutil
import sys
from dataclasses import dataclass
from datetime import datetime, timezone, timedelta
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.sax.saxutils import escape
from typing import Any, Dict, List, Optional, Tuple
from zoneinfo import ZoneInfo

SCHEMA_CONFIG = "raw-install-config/v1"
SCHEMA_LEDGER_EVENT = "raw-ledger-event/v1"
SCHEMA_ASSET = "raw-asset/v1"
SCHEMA_DELIVERY = "raw-delivery/v1"
SCHEMA_AUDIT = "raw-audit/v1"
SCHEMA_EVIDENCE_SEARCH_RESULT = "raw-evidence-search-result/v1"
SCHEMA_REPLAY_BUNDLE = "raw-replay-bundle/v1"


class RawOsError(Exception):
    pass


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z")


def parse_event_datetime(value: Optional[str]) -> Optional[datetime]:
    if value is None:
        return None
    raw = str(value).strip()
    if not raw:
        return None
    if re.match(r"^-?\d+(?:\.\d+)?$", raw):
        try:
            num = float(raw)
        except ValueError:
            return None
        # OpenClaw/Telegram metadata may expose JavaScript epoch milliseconds.
        # Smaller numeric values are treated as Unix seconds.
        if abs(num) > 10_000_000_000:
            num = num / 1000.0
        try:
            return datetime.fromtimestamp(num, tz=timezone.utc)
        except (OverflowError, OSError, ValueError):
            return None
    try:
        if raw.endswith("Z"):
            return datetime.fromisoformat(raw[:-1] + "+00:00")
        dt = datetime.fromisoformat(raw)
        if dt.tzinfo is None:
            return dt.replace(tzinfo=timezone.utc)
        return dt
    except ValueError:
        return None


def event_local_day(ts: Optional[str], tz_name: str) -> Optional[str]:
    dt = parse_event_datetime(ts)
    if not dt:
        return None
    try:
        tz = ZoneInfo(tz_name)
    except Exception:
        tz = timezone.utc
    return dt.astimezone(tz).date().isoformat()


def event_matches_anchor_day(ts: Optional[str], anchor_day: str, config: Dict[str, Any]) -> bool:
    """Return whether an event timestamp belongs to the requested local day.

    v0 intentionally uses event-time local calendar day rather than transcript
    mtime. If a timestamp cannot be parsed we keep the event instead of silently
    losing evidence; audit/incident layers can flag malformed inputs later.
    """
    day = event_local_day(ts, str((config.get("time") or {}).get("timezone") or "UTC"))
    return day is None or day == anchor_day


def read_json(path: Path) -> Dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, obj: Dict[str, Any]) -> None:
    ensure_dir(path.parent)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def append_jsonl(path: Path, obj: Dict[str, Any]) -> None:
    ensure_dir(path.parent)
    with path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(obj, ensure_ascii=False, sort_keys=True) + "\n")


def read_jsonl(path: Path) -> List[Dict[str, Any]]:
    if not path.exists():
        return []
    out: List[Dict[str, Any]] = []
    for i, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
        line = line.strip()
        if not line:
            continue
        try:
            out.append(json.loads(line))
        except json.JSONDecodeError as exc:
            raise RawOsError(f"invalid jsonl {path}:{i}: {exc}") from exc
    return out


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def protected_live_workspace_roots() -> List[Path]:
    raw = os.environ.get("RAW_OS_PROTECTED_WORKSPACE_ROOTS", "")
    roots: List[Path] = []
    for item in raw.split(os.pathsep):
        item = item.strip()
        if item:
            roots.append(Path(item).expanduser().resolve())
    return roots


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def stable_id(prefix: str, *parts: str, length: int = 16) -> str:
    h = hashlib.sha256("\x1f".join(parts).encode("utf-8")).hexdigest()[:length]
    return f"{prefix}_{h}"


def infer_kind(mime_type: Optional[str], path: Optional[Path] = None) -> str:
    mt = mime_type or ""
    if mt.startswith("image/"):
        return "image"
    if mt.startswith("audio/"):
        return "audio"
    if mt.startswith("video/"):
        return "video"
    if mt:
        return "file"
    if path:
        guessed, _ = mimetypes.guess_type(path.name)
        return infer_kind(guessed)
    return "file"


def config_errors(config: Dict[str, Any]) -> List[str]:
    errors: List[str] = []
    if config.get("schema") != SCHEMA_CONFIG:
        errors.append(f"schema must be {SCHEMA_CONFIG}")
    agent = config.get("agent") or {}
    if not agent.get("workspaceRoot"):
        errors.append("agent.workspaceRoot missing")
    if not agent.get("agentId"):
        errors.append("agent.agentId missing")
    if not config.get("mainlines"):
        errors.append("mainlines must not be empty")
    for i, ml in enumerate(config.get("mainlines", [])):
        if not ml.get("name"):
            errors.append(f"mainlines[{i}].name missing")
        if not ml.get("senderIds"):
            errors.append(f"mainlines[{i}].senderIds missing")
    storage = config.get("storage") or {}
    required_storage = ["ledgerDir", "assetStoreDir", "assetRegistryDir", "officialMdDir", "incidentDir", "stateDir"]
    optional_storage = ["memoryMdDir", "deliveryRegistryDir", "officialDocxDir"]
    for key in required_storage + optional_storage:
        value = storage.get(key)
        if not value:
            if key in required_storage:
                errors.append(f"storage.{key} missing")
            continue
        p = Path(str(value))
        if p.is_absolute():
            errors.append(f"storage.{key} must be relative to agent.workspaceRoot")
        if ".." in p.parts:
            errors.append(f"storage.{key} must not contain ..")
    return errors


@dataclass(frozen=True)
class RawOsContext:
    config_path: Path
    config: Dict[str, Any]
    root: Path
    storage: Dict[str, Path]

    @classmethod
    def load(cls, config_path: Path) -> "RawOsContext":
        config_path = config_path.resolve()
        config = read_json(config_path)
        errors = config_errors(config)
        if errors:
            raise RawOsError("invalid config: " + "; ".join(errors))
        root = Path(config["agent"]["workspaceRoot"]).expanduser().resolve()
        storage_config = dict(config.get("storage") or {})
        storage_config.setdefault("memoryMdDir", "./raw-memory-md")
        storage_config.setdefault("deliveryRegistryDir", "./delivery-registry")
        storage_config.setdefault("officialDocxDir", "./raw-docx")
        storage = {k: (root / v).resolve() for k, v in storage_config.items()}
        return cls(config_path=config_path, config=config, root=root, storage=storage)

    @property
    def agent_id(self) -> str:
        return self.config["agent"]["agentId"]

    def is_openclaw_workspace(self) -> bool:
        parts = self.root.parts
        return len(parts) >= 2 and parts[-1] == "workspace" and parts[-2].startswith(".openclaw")

    def assert_write_allowed(self) -> None:
        """Guard against accidentally writing into live OpenClaw workspaces.

        Raw OS development should default to temp/worktree targets. Writing into
        any live `.openclaw*/workspace` requires an explicit environment opt-in.
        Protected live workspace roots can add an additional dedicated opt-in
        for production baselines that should not be touched during experiments.
        """
        if self.root in protected_live_workspace_roots() and os.environ.get("RAW_OS_ALLOW_PROTECTED_WORKSPACE") != "1":
            raise RawOsError(
                "refusing to write to protected live workspace; set RAW_OS_ALLOW_PROTECTED_WORKSPACE=1 only for an intentional production operation"
            )
        if self.is_openclaw_workspace() and os.environ.get("RAW_OS_ALLOW_OPENCLAW_WORKSPACE") != "1":
            raise RawOsError(
                "refusing to write to live OpenClaw workspace; set RAW_OS_ALLOW_OPENCLAW_WORKSPACE=1 for an intentional deploy/test target"
            )

    def ensure_storage(self) -> None:
        self.assert_write_allowed()
        for key in ["ledgerDir", "assetStoreDir", "assetRegistryDir", "deliveryRegistryDir", "officialMdDir", "officialDocxDir", "memoryMdDir", "incidentDir", "stateDir"]:
            ensure_dir(self.storage[key])

    def ledger_path(self, anchor_day: str, mainline: str) -> Path:
        return self.storage["ledgerDir"] / anchor_day / f"{self.agent_id}__{mainline}.jsonl"

    def asset_registry_path(self, anchor_day: str) -> Path:
        return self.storage["assetRegistryDir"] / f"{anchor_day}.jsonl"

    def delivery_registry_path(self, anchor_day: str) -> Path:
        return self.storage["deliveryRegistryDir"] / f"{anchor_day}.jsonl"

    def official_md_path(self, anchor_day: str, mainline: str) -> Path:
        return self.storage["officialMdDir"] / f"{anchor_day}__{mainline}_raw.md"

    def official_docx_path(self, anchor_day: str, mainline: str) -> Path:
        return self.storage["officialDocxDir"] / f"{anchor_day}__{mainline}_raw.docx"

    def memory_md_path(self, anchor_day: str, mainline: str) -> Path:
        return self.storage["memoryMdDir"] / f"{anchor_day}__{mainline}_memory.md"

    def transcript_state_path(self, anchor_day: str, mainline: str) -> Path:
        return self.storage["stateDir"] / "transcript-ingest" / f"{anchor_day}__{mainline}.json"

    def tap_spool_state_path(self, anchor_day: str, mainline: str) -> Path:
        return self.storage["stateDir"] / "tap-spool-ingest" / f"{anchor_day}__{mainline}.json"

    def incident_dir(self) -> Path:
        return self.storage.get("incidentDir") or (self.storage["stateDir"] / "incidents")


def parse_kv_items(items: Optional[List[str]]) -> Dict[str, str]:
    out: Dict[str, str] = {}
    for item in items or []:
        if "=" not in item:
            raise RawOsError(f"metadata item must be key=value: {item}")
        k, v = item.split("=", 1)
        out[k.strip()] = v.strip()
    return out


def text_from_content_parts(parts: List[Dict[str, Any]]) -> str:
    texts: List[str] = []
    for part in parts:
        if isinstance(part, dict) and part.get("type") == "text" and part.get("text"):
            texts.append(str(part["text"]))
    return "\n".join(texts).strip()


def paths_from_attachment_refs(refs: List[Dict[str, Any]]) -> Tuple[List[Dict[str, Any]], List[str]]:
    local: List[Dict[str, Any]] = []
    failed: List[str] = []
    for i, ref in enumerate(refs):
        if not isinstance(ref, dict):
            failed.append(f"invalid-attachment-ref:{i}")
            continue
        status = ref.get("capture_status") or ref.get("status")
        local_path = ref.get("local_path") or ref.get("path")
        if local_path and status not in {"failed", "expired", "unsupported"}:
            local.append(ref)
        else:
            failed.append(str(ref.get("uri") or ref.get("file_id") or ref.get("name") or f"attachment:{i}"))
    return local, failed


def strip_openclaw_embedded_file_bodies(text: str) -> str:
    """Remove inline <file> bodies after media refs have been extracted.

    OpenClaw transcripts may inline uploaded text files for model context. Raw OS
    should not duplicate file bodies in raw-md once the asset is captured; it
    keeps the human message plus stable asset pointers instead.
    """
    cleaned = re.sub(r"\n?<file\b[^>]*>.*?</file>", "", text or "", flags=re.S | re.I)
    cleaned = re.sub(r"^\[media attached:[^\n]*\]\n?", "", cleaned, flags=re.M)
    return cleaned.strip()


def attachment_refs_from_openclaw_text(text: str) -> List[Dict[str, Any]]:
    """Extract OpenClaw inbound media markers from transcript text.

    OpenClaw renders inbound files into text parts as lines like:
    [media attached: /abs/path/file.md (text/markdown) | /abs/path/file.md]

    Raw OS treats those paths as evidence assets and copies them into the
    content-addressed asset store. Missing/unreadable paths are still returned
    as failed refs so the audit trail records the attempted attachment.
    """
    refs: List[Dict[str, Any]] = []
    seen = set()
    pattern = re.compile(r"\[media attached:\s*(?P<label>.*?)(?:\s*\((?P<mime>[^)]*)\))?\s*\|\s*(?P<path>[^\]]+)\]", re.I)
    for match in pattern.finditer(text or ""):
        raw_path = match.group("path").strip()
        label = match.group("label").strip()
        mime = (match.group("mime") or "").strip()
        key = raw_path or label
        if not key or key in seen:
            continue
        seen.add(key)
        ref: Dict[str, Any] = {"source": "openclaw-media-marker"}
        if raw_path.startswith("/") or raw_path.startswith("~"):
            ref["local_path"] = raw_path
        else:
            ref["uri"] = raw_path
            ref["status"] = "unsupported"
        if label:
            ref["name"] = Path(label.split(" (")[0]).name
        if mime:
            ref["mime_type"] = mime
        refs.append(ref)
    return refs


def sender_id_from_event(event: Dict[str, Any]) -> str:
    if event.get("sender_id"):
        return str(event["sender_id"])
    sender = event.get("sender")
    if isinstance(sender, dict) and sender.get("id"):
        return str(sender["id"])
    return "unknown-sender"


def extract_text_from_message_content(content: List[Dict[str, Any]]) -> str:
    texts: List[str] = []
    for part in content or []:
        if isinstance(part, dict) and part.get("type") == "text" and part.get("text"):
            texts.append(str(part["text"]))
    return "\n".join(texts).strip()


def parse_untrusted_metadata_blocks(text: str) -> Tuple[Dict[str, Any], Dict[str, Any], str]:
    """Parse OpenClaw user-visible metadata wrappers without trusting them.

    The returned dicts are treated as source metadata only. The cleaned text is
    the human message body after removing known wrapper blocks.
    """
    conversation: Dict[str, Any] = {}
    sender: Dict[str, Any] = {}

    def _load_after(label: str) -> Dict[str, Any]:
        pattern = re.compile(re.escape(label) + r":\s*```json\s*(\{.*?\})\s*```", re.S)
        m = pattern.search(text)
        if not m:
            return {}
        try:
            parsed = json.loads(m.group(1))
            return parsed if isinstance(parsed, dict) else {}
        except json.JSONDecodeError:
            return {}

    conversation = _load_after("Conversation info (untrusted metadata)")
    sender = _load_after("Sender (untrusted metadata)")
    cleaned = re.sub(r"Conversation info \(untrusted metadata\):\s*```json\s*\{.*?\}\s*```\s*", "", text, flags=re.S)
    cleaned = re.sub(r"Sender \(untrusted metadata\):\s*```json\s*\{.*?\}\s*```\s*", "", cleaned, flags=re.S)
    return conversation, sender, cleaned.strip()


def parse_openclaw_display_timestamp(value: Optional[str], fallback: Optional[str]) -> str:
    if not value:
        dt = parse_event_datetime(fallback)
        return dt.astimezone(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z") if dt else (fallback or utc_now_iso())
    dt = parse_event_datetime(value)
    if dt:
        return dt.astimezone(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")
    m = re.match(r"^[A-Za-z]{3}\s+(\d{4}-\d{2}-\d{2})\s+(\d{2}:\d{2})(?:\s+GMT([+-]\d+))?$", value.strip())
    if not m:
        return fallback or utc_now_iso()
    dt = datetime.fromisoformat(f"{m.group(1)}T{m.group(2)}:00")
    offset_hours = int(m.group(3) or "+0")
    aware = dt.replace(tzinfo=timezone(timedelta(hours=offset_hours)))
    return aware.astimezone(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")



CHANNEL_DISPLAY = {
    "telegram": "Telegram",
    "weixin": "微信",
    "wechat": "微信",
    "imessage": "iMessage",
    "bluebubbles": "iMessage",
    "whatsapp": "WhatsApp",
    "signal": "Signal",
    "discord": "Discord",
    "slack": "Slack",
    "feishu": "飞书",
    "lark": "飞书",
    "msteams": "Teams",
    "teams": "Teams",
    "matrix": "Matrix",
    "qq": "QQ",
    "qqbot": "QQ",
    "email": "Email",
    "himalaya": "Email",
}

PROVIDER_CHANNEL_ALIASES = {
    "telegram": "telegram",
    "openclaw-weixin": "weixin",
    "weixin": "weixin",
    "wechat": "weixin",
    "bluebubbles": "imessage",
    "imessage": "imessage",
    "whatsapp": "whatsapp",
    "signal": "signal",
    "discord": "discord",
    "slack": "slack",
    "feishu": "feishu",
    "lark": "feishu",
    "msteams": "msteams",
    "matrix": "matrix",
    "qqbot": "qq",
    "qq": "qq",
    "himalaya": "email",
    "email": "email",
}


def canonical_channel(value: Optional[str]) -> Optional[str]:
    if not value:
        return None
    raw = value.lower()
    if raw in PROVIDER_CHANNEL_ALIASES:
        return PROVIDER_CHANNEL_ALIASES[raw]
    if raw in CHANNEL_DISPLAY:
        return raw
    return raw if raw else None


def infer_channel_from_source(*, chat_id: str, message_id: str, provider: Optional[str] = None, default_channel: Optional[str] = None) -> str:
    raw = f"{chat_id} {message_id} {provider or ''}".lower()

    # 1. Structured chat_id prefix wins when present.
    if ":" in chat_id:
        prefixed = canonical_channel(chat_id.split(":", 1)[0])
        if prefixed:
            return prefixed

    # 2. Strong source patterns can override a generic/default provider.
    if "openclaw-weixin" in raw or "@im.wechat" in raw or "weixin" in raw or "wechat" in raw:
        return "weixin"

    # 3. Provider / plugin id.
    provider_channel = canonical_channel(provider)
    if provider_channel and provider_channel != "unknown":
        return provider_channel

    # 4. Remaining known source patterns.
    if "telegram:" in raw or "telegram_" in raw or "telegram" in raw:
        return "telegram"
    if "bluebubbles" in raw or "imessage" in raw:
        return "imessage"
    if "whatsapp" in raw:
        return "whatsapp"
    if "signal" in raw:
        return "signal"
    if "discord" in raw:
        return "discord"
    if "slack" in raw:
        return "slack"
    if "feishu" in raw or "lark" in raw:
        return "feishu"
    if "msteams" in raw or "teams" in raw:
        return "msteams"
    if "matrix" in raw:
        return "matrix"
    if "qqbot" in raw or "qq" in raw:
        return "qq"
    if "himalaya" in raw or "email" in raw:
        return "email"

    # Fallback must be unknown; never let a default channel pollute source truth.
    return "unknown"


def infer_provider_from_source(*, chat_id: str, message_id: str, default_provider: Optional[str], channel: str) -> str:
    raw = f"{chat_id} {message_id} {default_provider or ''}".lower()
    if "openclaw-weixin" in raw or "@im.wechat" in raw:
        return "openclaw-weixin"
    if default_provider and canonical_channel(default_provider) == channel and channel != "unknown":
        return default_provider
    if channel != "unknown":
        return channel
    return default_provider or "unknown"


def normalized_event_from_transcript_message(obj: Dict[str, Any], *, default_channel: str, default_provider: str, default_sender_id: Optional[str] = None, default_chat_id: Optional[str] = None, runtime_context: Optional[str] = None) -> Optional[Dict[str, Any]]:
    if obj.get("type") != "message":
        return None
    msg = obj.get("message") or {}
    role = msg.get("role")
    if role not in {"user", "assistant"}:
        return None
    content = msg.get("content") or []
    raw_text = extract_text_from_message_content(content)
    if not raw_text or raw_text.strip() == "[assistant turn failed before producing content]":
        return None

    if role == "user":
        conv, sender_meta, clean_text = parse_untrusted_metadata_blocks(raw_text)
        if runtime_context and not conv:
            ctx_conv, ctx_sender, _ctx_clean = parse_untrusted_metadata_blocks(runtime_context)
            conv = ctx_conv or conv
            sender_meta = ctx_sender or sender_meta
        attachment_refs = attachment_refs_from_openclaw_text(raw_text)
        clean_text = strip_openclaw_embedded_file_bodies(clean_text)
        if not clean_text and not attachment_refs:
            return None
        chat_id = str(conv.get("chat_id") or default_chat_id or "unknown-chat")
        message_id = str(conv.get("message_id") or obj.get("id") or "unknown-message")
        sender_id = str(conv.get("sender_id") or sender_meta.get("id") or default_sender_id or "unknown-sender")
        channel = infer_channel_from_source(chat_id=chat_id, message_id=message_id, provider=conv.get("provider") or default_provider, default_channel=default_channel)
        provider = infer_provider_from_source(chat_id=chat_id, message_id=message_id, default_provider=default_provider, channel=channel)
        timestamp = str(obj.get("timestamp") or parse_openclaw_display_timestamp(conv.get("timestamp"), obj.get("timestamp")))
        speaker = "user"
        sender_label = "你"
        sender_label_meta = sender_meta.get("label") or conv.get("sender")
        event_id = f"{channel}_{sender_id}_{message_id}"
    else:
        clean_text = raw_text.strip()
        chat_id = str(default_chat_id or "unknown-chat")
        message_id = str(obj.get("id") or "unknown-message")
        sender_id = "assistant"
        channel = chat_id.split(":", 1)[0] if ":" in chat_id else default_channel
        provider = default_provider or channel
        timestamp = parse_openclaw_display_timestamp(None, obj.get("timestamp"))
        speaker = "assistant"
        sender_label = "我"
        sender_label_meta = "assistant"
        event_id = f"{channel}_assistant_{message_id}"
        sender_meta = {}

    if role != "user":
        attachment_refs = []

    return {
        "event_id": event_id,
        "timestamp": timestamp,
        "channel": channel,
        "provider": provider,
        "chat_type": "direct",
        "chat_id": chat_id,
        "message_id": message_id,
        "session_key": f"transcript:{Path(str(obj.get('session_file', 'unknown'))).stem}",
        "sender": {"id": sender_id, "label": sender_label_meta},
        "sender_label": sender_label,
        "speaker": speaker,
        "content_parts": [{"type": "text", "text": clean_text}],
        "attachment_refs": attachment_refs,
        "ingest_metadata": {
            "source": "openclaw-session-transcript",
            "transcript_record_id": obj.get("id"),
            "parent_id": obj.get("parentId"),
            "metadata_untrusted": role == "user",
            "sender_username": sender_meta.get("username"),
        },
    }


def store_asset(ctx: RawOsContext, *, source_path: Path, anchor_day: str) -> Tuple[Dict[str, Any], Path]:
    if not source_path.exists() or not source_path.is_file():
        raise RawOsError(f"attachment path not found: {source_path}")
    digest = sha256_file(source_path)
    guessed_mime, _ = mimetypes.guess_type(source_path.name)
    ext = source_path.suffix or mimetypes.guess_extension(guessed_mime or "") or ""
    dest = ctx.storage["assetStoreDir"] / "by-sha256" / digest[:2] / digest[2:4] / f"{digest}{ext}"
    ensure_dir(dest.parent)
    if not dest.exists():
        shutil.copy2(source_path, dest)
    rel_dest = os.path.relpath(dest, ctx.root)
    info = {
        "sha256": digest,
        "byte_size": source_path.stat().st_size,
        "mime_type": guessed_mime or "application/octet-stream",
        "kind": infer_kind(guessed_mime, source_path),
        "original_filename": source_path.name,
        "storage_uri": rel_dest,
    }
    return info, dest


def build_asset_record(
    ctx: RawOsContext,
    *,
    anchor_day: str,
    event_id: str,
    mainline: str,
    source_ref: Dict[str, Any],
    source_path: Optional[Path],
    failed_ref: Optional[str],
    failure_reason: Optional[str],
    attachment_index: int,
    caption: Optional[str],
    surrounding_text: Optional[str],
) -> Dict[str, Any]:
    now = utc_now_iso()
    source_ref = dict(source_ref)
    source_ref["attachment_index"] = attachment_index

    if source_path:
        stored, _dest = store_asset(ctx, source_path=source_path, anchor_day=anchor_day)
        if source_ref.get("mime_type"):
            stored["mime_type"] = str(source_ref["mime_type"])
            stored["kind"] = infer_kind(str(source_ref["mime_type"]), source_path)
        if source_ref.get("name"):
            stored["original_filename"] = str(source_ref["name"])
        asset_id = stable_id("asset", event_id, stored["sha256"], str(attachment_index))
        capture = {"status": "captured", "captured_at": now, "method": "local-copy", "error": None}
        base = stored
    else:
        ref = failed_ref or f"missing:{attachment_index}"
        asset_id = stable_id("asset", event_id, ref, str(attachment_index))
        capture = {"status": "failed", "captured_at": None, "method": "unavailable", "error": failure_reason or "capture failed"}
        base = {
            "sha256": None,
            "byte_size": None,
            "mime_type": None,
            "kind": "file",
            "original_filename": failed_ref,
            "storage_uri": None,
        }

    return {
        "schema": SCHEMA_ASSET,
        "asset_id": asset_id,
        "source_event_id": event_id,
        "source_ref": source_ref,
        "mainline": mainline,
        "original_filename": base["original_filename"],
        "mime_type": base["mime_type"],
        "byte_size": base["byte_size"],
        "sha256": base["sha256"],
        "kind": base["kind"],
        "storage_uri": base["storage_uri"],
        "capture": capture,
        "context": {"caption": caption, "surrounding_text": surrounding_text, "sender_label": None},
        "derived": {
            "ocr": {"status": "pending" if base["kind"] == "image" and source_path else "not_applicable", "text": None, "model": None, "updated_at": None},
            "transcription": {"status": "pending" if base["kind"] == "audio" and source_path else "not_applicable", "text": None, "model": None, "updated_at": None},
            "summary": {"status": "pending" if source_path else "not_available", "text": None, "model": None, "updated_at": None},
        },
        "retention": {"policy": "default", "redacted": False, "redaction_reason": None},
        "audit": {"created_at": now, "updated_at": now},
    }


def build_delivery_record(
    *,
    anchor_day: str,
    event_id: str,
    mainline: str,
    source_ref: Dict[str, Any],
    delivery: Dict[str, Any],
    now: str,
) -> Dict[str, Any]:
    status = str(delivery.get("status") or ("failed" if delivery.get("error") else "sent"))
    target = str(delivery.get("target") or source_ref.get("chat_id") or "unknown-target")
    channel_message_id = delivery.get("channel_message_id") or delivery.get("message_id") or source_ref.get("message_id")
    delivery_id = str(delivery.get("delivery_id") or stable_id("delivery", anchor_day, mainline, event_id, target, status, str(channel_message_id)))
    return {
        "schema": SCHEMA_DELIVERY,
        "delivery_id": delivery_id,
        "anchor_day": anchor_day,
        "mainline": mainline,
        "source_event_id": event_id,
        "channel": delivery.get("channel") or source_ref.get("channel"),
        "provider": delivery.get("provider") or source_ref.get("provider"),
        "target": target,
        "chat_id": delivery.get("chat_id") or source_ref.get("chat_id"),
        "thread_id": delivery.get("thread_id"),
        "channel_message_id": channel_message_id,
        "status": status,
        "error": delivery.get("error"),
        "fallback_reason": delivery.get("fallback_reason"),
        "sent_at": delivery.get("sent_at") or (now if status == "sent" else None),
        "acknowledged_at": delivery.get("acknowledged_at"),
        "metadata": delivery.get("metadata") or {},
        "source_ref": source_ref,
        "created_at": now,
    }


def ingest_event(
    ctx: RawOsContext,
    *,
    anchor_day: str,
    mainline: str,
    text: str,
    speaker: str,
    sender_id: str,
    sender_label: Optional[str],
    channel: str,
    provider: str,
    chat_id: str,
    message_id: str,
    session_key: str,
    attachments: List[Any],
    failed_assets: List[str],
    failure_reason: Optional[str],
    metadata: Dict[str, Any],
    event_ts: Optional[str] = None,
    upstream_event_id: Optional[str] = None,
    reply_to: Optional[Dict[str, Any]] = None,
    content_parts: Optional[List[Dict[str, Any]]] = None,
    delivery: Optional[Dict[str, Any]] = None,
    event_type: str = "message",
) -> Dict[str, Any]:
    ctx.ensure_storage()
    now = utc_now_iso()
    ts = event_ts or now
    source_ref = {"channel": channel, "provider": provider, "chat_id": chat_id, "message_id": message_id, "upstream_event_id": upstream_event_id}
    attachment_fingerprint = "|".join([str(p) for p in attachments] + list(failed_assets))
    event_id = upstream_event_id or stable_id("evt", anchor_day, mainline, channel, chat_id, message_id, text, attachment_fingerprint)
    ledger_path = ctx.ledger_path(anchor_day, mainline)
    existing_ids = {item.get("event_id") for item in read_jsonl(ledger_path)}
    if event_id in existing_ids:
        return {"ok": True, "skipped": True, "reason": "duplicate_event_id", "event_id": event_id, "ledger": str(ledger_path), "assets": 0}

    asset_records: List[Dict[str, Any]] = []
    idx = 0
    for attachment in attachments:
        if isinstance(attachment, dict):
            path = Path(str(attachment.get("local_path") or attachment.get("path")))
            source_ref_for_asset = dict(source_ref)
            for key in ["name", "mime_type", "kind", "source"]:
                if attachment.get(key):
                    source_ref_for_asset[key] = attachment.get(key)
        else:
            path = Path(attachment)
            source_ref_for_asset = source_ref
        asset_records.append(build_asset_record(
            ctx,
            anchor_day=anchor_day,
            event_id=event_id,
            mainline=mainline,
            source_ref=source_ref_for_asset,
            source_path=path.expanduser().resolve(),
            failed_ref=None,
            failure_reason=None,
            attachment_index=idx,
            caption=text or None,
            surrounding_text=text or None,
        ))
        idx += 1
    for ref in failed_assets:
        asset_records.append(build_asset_record(
            ctx,
            anchor_day=anchor_day,
            event_id=event_id,
            mainline=mainline,
            source_ref=source_ref,
            source_path=None,
            failed_ref=ref,
            failure_reason=failure_reason,
            attachment_index=idx,
            caption=text or None,
            surrounding_text=text or None,
        ))
        idx += 1

    registry_path = ctx.asset_registry_path(anchor_day)
    for record in asset_records:
        append_jsonl(registry_path, record)

    asset_pointers = []
    rel_registry = os.path.relpath(registry_path, ctx.root)
    for record in asset_records:
        asset_pointers.append({
            "asset_id": record["asset_id"],
            "kind": record["kind"],
            "mime_type": record["mime_type"],
            "sha256": record["sha256"],
            "capture_status": record["capture"]["status"],
            "registry_ref": f"{rel_registry}#{record['asset_id']}",
        })

    delivery_records: List[Dict[str, Any]] = []
    delivery_pointers: List[Dict[str, Any]] = []
    delivery_registry_path = ctx.delivery_registry_path(anchor_day)
    if delivery:
        record = build_delivery_record(anchor_day=anchor_day, event_id=event_id, mainline=mainline, source_ref=source_ref, delivery=delivery, now=now)
        append_jsonl(delivery_registry_path, record)
        delivery_records.append(record)
        rel_delivery_registry = os.path.relpath(delivery_registry_path, ctx.root)
        delivery_pointers.append({
            "delivery_id": record["delivery_id"],
            "status": record["status"],
            "target": record["target"],
            "channel_message_id": record.get("channel_message_id"),
            "registry_ref": f"{rel_delivery_registry}#{record['delivery_id']}",
        })

    event = {
        "schema": SCHEMA_LEDGER_EVENT,
        "event_id": event_id,
        "event_type": event_type,
        "ts": ts,
        "anchor_day": anchor_day,
        "agent_id": ctx.agent_id,
        "session_key": session_key,
        "channel": channel,
        "provider": provider,
        "chat_type": metadata.get("chat_type", "direct"),
        "chat_id": chat_id,
        "target_mainline": mainline,
        "speaker": speaker,
        "speaker_label": sender_label or ("你" if speaker == "user" else "我"),
        "sender_id": sender_id,
        "source_ref": source_ref,
        "reply_to": reply_to,
        "content_parts": content_parts if content_parts is not None else ([{"type": "text", "text": text}] if text else []),
        "raw_text": text,
        "clean_text": text.strip(),
        "assets": asset_pointers,
        "deliveries": delivery_pointers,
        "delivery": delivery if delivery else None,
        "classification": {"is_noise": False, "noise_kind": None, "reason": None, "classifier_version": "v0"},
        "metadata": metadata,
        "created_at": now,
    }
    append_jsonl(ledger_path, event)
    return {"ok": True, "skipped": False, "event_id": event_id, "ledger": str(ledger_path), "asset_registry": str(registry_path), "delivery_registry": str(delivery_registry_path), "assets": len(asset_records), "deliveries": len(delivery_records)}


def load_asset_registry(ctx: RawOsContext, anchor_day: str) -> Dict[str, Dict[str, Any]]:
    records = read_jsonl(ctx.asset_registry_path(anchor_day))
    return {r.get("asset_id"): r for r in records if r.get("asset_id")}


def load_delivery_registry(ctx: RawOsContext, anchor_day: str) -> Dict[str, Dict[str, Any]]:
    records = read_jsonl(ctx.delivery_registry_path(anchor_day))
    return {r.get("delivery_id"): r for r in records if r.get("delivery_id")}



def channel_display_label(channel: Optional[str], provider: Optional[str] = None) -> str:
    canonical = canonical_channel(channel) or canonical_channel(provider) or "unknown"
    return CHANNEL_DISPLAY.get(canonical, canonical if canonical != "unknown" else "unknown")


def render_speaker_label(event: Dict[str, Any], *, include_assistant_source: bool = False) -> str:
    speaker = event.get("speaker_label") or event.get("speaker") or "?"
    channel = channel_display_label(event.get("channel"), event.get("provider"))
    if event.get("speaker") == "assistant" and not include_assistant_source:
        return str(speaker)
    if channel == "unknown":
        return str(speaker)
    return f"{speaker}（{channel}）"


def find_asset_record(ctx: RawOsContext, *, asset_id: str, anchor_day: Optional[str] = None) -> Tuple[Dict[str, Any], Path]:
    registry_paths = [ctx.asset_registry_path(anchor_day)] if anchor_day else sorted(ctx.storage["assetRegistryDir"].glob("*.jsonl"), reverse=True)
    for registry_path in registry_paths:
        if not registry_path.exists():
            continue
        for record in read_jsonl(registry_path):
            if record.get("asset_id") == asset_id:
                return record, registry_path
    raise RawOsError(f"asset not found: {asset_id}")


def resolve_asset_path(ctx: RawOsContext, record: Dict[str, Any]) -> Path:
    uri = record.get("storage_uri")
    if not uri:
        raise RawOsError(f"asset has no storage_uri: {record.get('asset_id')}")
    path = (ctx.root / str(uri)).resolve()
    try:
        path.relative_to(ctx.root)
    except ValueError as exc:
        raise RawOsError(f"asset path escapes workspace: {uri}") from exc
    if not path.exists() or not path.is_file():
        raise RawOsError(f"asset file missing: {path}")
    expected = record.get("sha256")
    if expected and sha256_file(path) != expected:
        raise RawOsError(f"asset hash mismatch: {record.get('asset_id')}")
    return path


def export_asset(ctx: RawOsContext, *, asset_id: str, out_dir: Path, anchor_day: Optional[str] = None) -> Dict[str, Any]:
    record, registry_path = find_asset_record(ctx, asset_id=asset_id, anchor_day=anchor_day)
    source = resolve_asset_path(ctx, record)
    ensure_dir(out_dir)
    filename = record.get("original_filename") or source.name
    dest = out_dir / Path(str(filename)).name
    if dest.exists():
        stem, suffix = dest.stem, dest.suffix
        dest = out_dir / f"{stem}__{asset_id}{suffix}"
    shutil.copy2(source, dest)
    return {"ok": True, "asset_id": asset_id, "source": str(source), "output": str(dest), "registry": str(registry_path), "sha256": record.get("sha256")}


ASSET_ID_RE = re.compile(r"\basset_[0-9a-fA-F]{8,64}\b")


def asset_id_from_natural_language(text: str) -> Optional[str]:
    match = ASSET_ID_RE.search(text or "")
    return match.group(0) if match else None

def ingest_normalized_event(ctx: RawOsContext, *, anchor_day: str, mainline: str, event: Dict[str, Any]) -> Dict[str, Any]:
    """Ingest a Raw OS normalized event document.

    Expected v0 shape is intentionally loose while adapters mature:
    - content_parts: [{type: text, text: ...}, ...]
    - attachment_refs: [{local_path/path: ...}, {status: failed, uri: ...}, ...]
    - source_ref or top-level channel/provider/chat_id/message_id fields
    """
    source_ref = event.get("source_ref") or {}
    content_parts = event.get("content_parts") or []
    if not isinstance(content_parts, list):
        raise RawOsError("normalized event content_parts must be a list")
    attachment_refs = event.get("attachment_refs") or event.get("attachments") or []
    if not isinstance(attachment_refs, list):
        raise RawOsError("normalized event attachment_refs must be a list")
    attachments, failed_assets = paths_from_attachment_refs(attachment_refs)
    failure_reasons = [str(ref.get("capture_error") or ref.get("error")) for ref in attachment_refs if isinstance(ref, dict) and (ref.get("capture_error") or ref.get("error"))]
    text = str(event.get("clean_text") or event.get("raw_text") or text_from_content_parts(content_parts))
    # Runtime hooks may expose OpenClaw's AI-facing inbound metadata prefix
    # (Conversation info / Sender info). Raw OS should keep the metadata in
    # structured fields, not render it as human message text.
    _conv_meta, _sender_meta, stripped_text = parse_untrusted_metadata_blocks(text)
    if stripped_text != text:
        text = stripped_text
        content_parts = [{"type": "text", "text": text}] if text else []
    channel = str(event.get("channel") or source_ref.get("channel") or source_ref.get("provider") or "unknown")
    provider = str(event.get("provider") or source_ref.get("provider") or channel)
    chat_id = str(event.get("chat_id") or source_ref.get("chat_id") or "unknown-chat")
    message_id = str(event.get("message_id") or source_ref.get("message_id") or event.get("event_id") or "unknown-message")
    return ingest_event(
        ctx,
        anchor_day=anchor_day,
        mainline=mainline,
        text=text,
        speaker=str(event.get("speaker") or "user"),
        sender_id=sender_id_from_event(event),
        sender_label=event.get("sender_label"),
        channel=channel,
        provider=provider,
        chat_id=chat_id,
        message_id=message_id,
        session_key=str(event.get("session_key") or "raw-os:normalized-event"),
        attachments=attachments,
        failed_assets=failed_assets,
        failure_reason="; ".join(failure_reasons) if failure_reasons else None,
        metadata=event.get("ingest_metadata") or event.get("metadata") or {},
        event_ts=event.get("timestamp") or event.get("ts"),
        upstream_event_id=event.get("event_id"),
        reply_to=event.get("reply_to") or event.get("thread"),
        content_parts=content_parts,
        delivery=event.get("delivery") if isinstance(event.get("delivery"), dict) else None,
        event_type=str(event.get("event_type") or "message"),
    )


def event_render_sort_key(event: Dict[str, Any]) -> Tuple[float, str]:
    dt = parse_event_datetime(event.get("ts"))
    if dt:
        return (dt.timestamp(), str(event.get("event_id") or ""))
    return (float("inf"), str(event.get("event_id") or ""))


def sort_events_for_render(events: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    return sorted(events, key=event_render_sort_key)


def render_day(ctx: RawOsContext, *, anchor_day: str, mainline: str) -> Dict[str, Any]:
    ctx.assert_write_allowed()
    events = sort_events_for_render(read_jsonl(ctx.ledger_path(anchor_day, mainline)))
    assets = load_asset_registry(ctx, anchor_day)
    deliveries = load_delivery_registry(ctx, anchor_day)
    output = ctx.official_md_path(anchor_day, mainline)
    ensure_dir(output.parent)

    lines: List[str] = [
        "---",
        "schema: raw-md/v1",
        "kind: official-raw",
        f"anchor_day: {anchor_day}",
        f"timezone: {ctx.config.get('time', {}).get('timezone', 'UTC')}",
        f"agent_id: {ctx.agent_id}",
        f"target_mainline: {mainline}",
        "source_of_truth: raw-ledger",
        f"content_status: {'complete' if events else 'partial'}",
        f"coverage_status: {'verified' if events else 'gap-detected'}",
        "delivery_status: pending",
        "---",
        "",
        "# Raw Report",
        "",
        "## Conversation",
        "",
    ]

    for event in events:
        dt = parse_event_datetime(event.get("ts"))
        if dt:
            tz_name = str(ctx.config.get("time", {}).get("timezone", "UTC"))
            try:
                tz = ZoneInfo(tz_name)
            except Exception:
                tz = timezone.utc
            ts = dt.astimezone(tz).strftime("%Y-%m-%d %H:%M:%S %Z")
        else:
            ts = str(event.get("ts", "")).replace("T", " ").replace("Z", " UTC")
        speaker = render_speaker_label(event)
        text = official_render_text_for_event(event)
        lines.append(f"[{ts}] {speaker}：{text}")
        delivery_ptrs = event.get("deliveries") or []
        if delivery_ptrs:
            lines.append("投递：")
            for ptr in delivery_ptrs:
                record = deliveries.get(ptr.get("delivery_id"), {})
                status = record.get("status") or ptr.get("status") or "unknown"
                target = record.get("target") or ptr.get("target") or "unknown-target"
                channel_message_id = record.get("channel_message_id") or ptr.get("channel_message_id")
                lines.append(f"- status: {status} · target: {target}")
                lines.append(f"  delivery_id: {ptr.get('delivery_id')}")
                if channel_message_id:
                    lines.append(f"  channel_message_id: {channel_message_id}")
                if record.get("error"):
                    lines.append(f"  error: {record.get('error')}")
        pointers = event.get("assets") or []
        if pointers:
            lines.append("附件：")
            for ptr in pointers:
                record = assets.get(ptr.get("asset_id"), {})
                capture = (record.get("capture") or {}).get("status") or ptr.get("capture_status")
                storage_uri = record.get("storage_uri")
                original = record.get("original_filename") or ptr.get("asset_id")
                lines.append(f"- {ptr.get('mime_type') or record.get('mime_type') or 'unknown'} · {original}")
                lines.append(f"  asset_id: {ptr.get('asset_id')}")
                lines.append(f"  sha256: {ptr.get('sha256') or record.get('sha256')}")
                lines.append(f"  path: {storage_uri}")
                lines.append(f"  retrieve: 让 agent 发送 {ptr.get('asset_id')}")
                lines.append(f"  capture: {capture}")
                if capture != "captured":
                    err = ((record.get("capture") or {}).get("error")) or "unknown"
                    lines.append(f"  capture_error: {err}")
        lines.append("")

    output.write_text("\n".join(lines), encoding="utf-8")
    return {"ok": True, "output": str(output), "events": len(events)}


def _write_official_docx_ooxml_fallback(rows: List[Tuple[str, str, str, str]], output: Path, title: str) -> None:
    # Dependency-free .docx writer for cron/runtime Python environments without python-docx.
    def run(text: str, color: str) -> str:
        body = []
        for i, chunk in enumerate(str(text).split("\n")):
            if i:
                body.append("<w:br/>")
            body.append(f'<w:t xml:space="preserve">{escape(chunk)}</w:t>')
        return (
            "<w:r><w:rPr>"
            "<w:rFonts w:ascii=\"PingFang SC\" w:hAnsi=\"PingFang SC\" w:eastAsia=\"PingFang SC\"/>"
            f"<w:color w:val=\"{color}\"/>"
            "</w:rPr>" + "".join(body) + "</w:r>"
        )

    paragraphs = [
        "<w:p><w:r><w:rPr><w:b/><w:sz w:val=\"32\"/>"
        "<w:rFonts w:ascii=\"PingFang SC\" w:hAnsi=\"PingFang SC\" w:eastAsia=\"PingFang SC\"/>"
        f"</w:rPr><w:t>{escape(title)}</w:t></w:r></w:p>"
    ]
    for ts, speaker, text, color in rows:
        line = f"[{ts}] {speaker}：{text}" if ts else f"{speaker}：{text}"
        paragraphs.append(f"<w:p>{run(line, color)}</w:p>")

    document = """<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>
<w:document xmlns:w=\"http://schemas.openxmlformats.org/wordprocessingml/2006/main\"><w:body>""" + "".join(paragraphs) + """<w:sectPr><w:pgSz w:w=\"11906\" w:h=\"16838\"/><w:pgMar w:top=\"1440\" w:right=\"1440\" w:bottom=\"1440\" w:left=\"1440\"/></w:sectPr></w:body></w:document>"""
    styles = """<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>
<w:styles xmlns:w=\"http://schemas.openxmlformats.org/wordprocessingml/2006/main\"><w:style w:type=\"paragraph\" w:default=\"1\" w:styleId=\"Normal\"><w:name w:val=\"Normal\"/><w:rPr><w:rFonts w:ascii=\"PingFang SC\" w:hAnsi=\"PingFang SC\" w:eastAsia=\"PingFang SC\"/></w:rPr></w:style></w:styles>"""
    font_table = """<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>
<w:fonts xmlns:w=\"http://schemas.openxmlformats.org/wordprocessingml/2006/main\"><w:font w:name=\"PingFang SC\"><w:charset w:val=\"86\"/><w:family w:val=\"swiss\"/></w:font></w:fonts>"""
    content_types = """<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>
<Types xmlns=\"http://schemas.openxmlformats.org/package/2006/content-types\"><Default Extension=\"rels\" ContentType=\"application/vnd.openxmlformats-package.relationships+xml\"/><Default Extension=\"xml\" ContentType=\"application/xml\"/><Override PartName=\"/word/document.xml\" ContentType=\"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml\"/><Override PartName=\"/word/styles.xml\" ContentType=\"application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml\"/><Override PartName=\"/word/fontTable.xml\" ContentType=\"application/vnd.openxmlformats-officedocument.wordprocessingml.fontTable+xml\"/><Override PartName=\"/docProps/core.xml\" ContentType=\"application/vnd.openxmlformats-package.core-properties+xml\"/><Override PartName=\"/docProps/app.xml\" ContentType=\"application/vnd.openxmlformats-officedocument.extended-properties+xml\"/></Types>"""
    rels = """<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>
<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\"><Relationship Id=\"rId1\" Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument\" Target=\"word/document.xml\"/><Relationship Id=\"rId2\" Type=\"http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties\" Target=\"docProps/core.xml\"/><Relationship Id=\"rId3\" Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties\" Target=\"docProps/app.xml\"/></Relationships>"""
    doc_rels = """<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>
<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\"/>"""
    core = f"""<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>
<cp:coreProperties xmlns:cp=\"http://schemas.openxmlformats.org/package/2006/metadata/core-properties\" xmlns:dc=\"http://purl.org/dc/elements/1.1/\"><dc:title>{escape(title)}</dc:title></cp:coreProperties>"""
    app = """<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>
<Properties xmlns=\"http://schemas.openxmlformats.org/officeDocument/2006/extended-properties\"><Application>Raw OS</Application></Properties>"""

    output.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(output, "w", ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", document)
        z.writestr("word/styles.xml", styles)
        z.writestr("word/fontTable.xml", font_table)
        z.writestr("word/_rels/document.xml.rels", doc_rels)
        z.writestr("docProps/core.xml", core)
        z.writestr("docProps/app.xml", app)


def render_official_docx(ctx: RawOsContext, *, anchor_day: str, mainline: str) -> Dict[str, Any]:
    ctx.assert_write_allowed()
    events = sort_events_for_render(read_jsonl(ctx.ledger_path(anchor_day, mainline)))
    output = ctx.official_docx_path(anchor_day, mainline)
    ensure_dir(output.parent)

    tz_name = str(ctx.config.get("time", {}).get("timezone", "UTC"))
    try:
        tz = ZoneInfo(tz_name)
    except Exception:
        tz = timezone.utc

    rows: List[Tuple[str, str, str, str]] = []
    for event in events:
        dt = parse_event_datetime(event.get("ts"))
        if dt:
            ts = dt.astimezone(tz).strftime("%Y-%m-%d %H:%M:%S %Z")
        else:
            ts = str(event.get("ts", "")).replace("T", " ").replace("Z", " UTC")
        speaker = render_speaker_label(event)
        text = official_render_text_for_event(event)
        color = "C00000" if event.get("speaker") == "user" else "000000"
        rows.append((ts, speaker, text, color))
        delivery_ptrs = event.get("deliveries") or []
        if delivery_ptrs:
            rows.append(("", "投递", "; ".join(str(ptr.get("delivery_id")) for ptr in delivery_ptrs), "666666"))
        pointers = event.get("assets") or []
        if pointers:
            rows.append(("", "附件", "; ".join(str(ptr.get("asset_id")) for ptr in pointers), "666666"))

    title = f"Raw Report {anchor_day} {mainline}"
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import RGBColor
    except Exception:
        _write_official_docx_ooxml_fallback(rows, output, title)
        return {"ok": True, "output": str(output), "events": len(events), "renderer": "ooxml-fallback"}

    doc = Document()
    doc.core_properties.title = title
    normal = doc.styles["Normal"]
    normal.font.name = "PingFang SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")

    doc.add_heading("Raw Report", level=1)
    doc.add_paragraph(f"anchor_day: {anchor_day} · mainline: {mainline} · timezone: {tz_name}")

    def add_run(paragraph, text: str, color: str) -> None:
        parts = str(text).split("\n")
        for idx, part in enumerate(parts):
            if idx:
                paragraph.add_run().add_break()
            run = paragraph.add_run(part)
            run.font.name = "PingFang SC"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
            run.font.color.rgb = RGBColor.from_string(color)

    for ts, speaker, text, color in rows:
        p = doc.add_paragraph()
        if ts:
            add_run(p, f"[{ts}] {speaker}：{text}", color)
        else:
            add_run(p, f"{speaker}：{text}", color)

    doc.save(str(output))
    return {"ok": True, "output": str(output), "events": len(events), "renderer": "python-docx"}


def is_automation_trigger_text(text: str) -> bool:
    stripped = (text or "").lstrip()
    return stripped.startswith("[cron:") or stripped.startswith("[heartbeat:") or "Current time:" in stripped and "Use the message tool" in stripped


def automation_trigger_title(text: str) -> str:
    first = (text or "").splitlines()[0] if text else "自动任务"
    m = re.match(r"\[cron:[^\s\]]+\s+([^\]]+)\]", first)
    if m:
        return m.group(1).strip()
    if first.startswith("[heartbeat:"):
        return "heartbeat"
    return "自动任务"


def compact_assistant_for_memory(text: str, *, max_chars: int = 1600) -> str:
    text = (text or "").strip()
    if len(text) <= max_chars:
        return text
    lines = [line.rstrip() for line in text.splitlines()]
    kept: List[str] = []
    budget = max_chars
    for line in lines:
        stripped = line.strip()
        if not stripped:
            if kept and kept[-1] != "":
                kept.append("")
            continue
        important = (
            stripped.startswith("#")
            or stripped.startswith("##")
            or stripped.startswith("###")
            or stripped.startswith("- ")
            or stripped.startswith("1.")
            or "一句话" in stripped
            or "结论" in stripped
            or "建议" in stripped
            or "判断" in stripped
            or "问题" in stripped
        )
        if important or not kept:
            add = line[: min(len(line), 220)]
            if len("\n".join(kept + [add])) > budget:
                break
            kept.append(add)
    compact = "\n".join(kept).strip()
    if not compact:
        compact = text[:max_chars].rstrip()
    return compact + "\n[长回答已折叠；完整内容见 raw-md / raw-ledger]"



def official_render_text_for_event(event: Dict[str, Any]) -> str:
    """Return human-facing official raw text while preserving full text in ledger.

    Automation prompts are operational instructions, not conversation content.
    Keep the event visible, but fold the prompt body so official raw/docx stays readable.
    """
    text = event.get("clean_text") or event.get("raw_text") or ""
    if event.get("speaker") == "user" and is_automation_trigger_text(text):
        return f"[自动任务: {automation_trigger_title(text)}]"
    return text

def project_text_for_memory(event: Dict[str, Any]) -> Tuple[Optional[str], bool]:
    text = (event.get("clean_text") or event.get("raw_text") or "").strip()
    if not text or text == "NO_REPLY" or text == "HEARTBEAT_OK":
        return None, False
    if event.get("speaker") == "user" and is_automation_trigger_text(text):
        return f"[自动任务: {automation_trigger_title(text)}]", True
    if event.get("speaker") == "assistant":
        return compact_assistant_for_memory(text), False
    return text, False


def render_memory_day(ctx: RawOsContext, *, anchor_day: str, mainline: str) -> Dict[str, Any]:
    ctx.assert_write_allowed()
    events = sort_events_for_render(read_jsonl(ctx.ledger_path(anchor_day, mainline)))
    assets = load_asset_registry(ctx, anchor_day)
    output = ctx.memory_md_path(anchor_day, mainline)
    ensure_dir(output.parent)

    lines: List[str] = [
        f"# Raw Memory Projection — {anchor_day} / {mainline}",
        "",
    ]
    projected = 0
    folded_automation = 0
    for event in events:
        projected_text, folded = project_text_for_memory(event)
        pointers = event.get("assets") or []
        if projected_text is None and not pointers:
            continue
        if folded:
            folded_automation += 1
        dt = parse_event_datetime(event.get("ts"))
        if dt:
            tz_name = str(ctx.config.get("time", {}).get("timezone", "UTC"))
            try:
                tz = ZoneInfo(tz_name)
            except Exception:
                tz = timezone.utc
            ts = dt.astimezone(tz).strftime("%Y-%m-%d %H:%M:%S %Z")
        else:
            ts = str(event.get("ts", "")).replace("T", " ").replace("Z", " UTC")
        speaker = render_speaker_label(event)
        if projected_text:
            lines.append(f"[{ts}] {speaker}：{projected_text}")
            projected += 1
        for ptr in pointers:
            record = assets.get(ptr.get("asset_id"), {})
            original = record.get("original_filename") or ptr.get("asset_id")
            mime = ptr.get("mime_type") or record.get("mime_type") or "unknown"
            lines.append(f"[附件: {original}, {mime}]")
        lines.append("")

    output.write_text("\n".join(lines), encoding="utf-8")
    collection_path = update_memory_collection_manifest(ctx, anchor_day=anchor_day, mainline=mainline, memory_md=output)
    return {"ok": True, "output": str(output), "events": len(events), "projected_events": projected, "folded_automation": folded_automation, "memory_collection": str(collection_path)}


def text_snippet(text: str, *, max_chars: int = 240) -> str:
    compact = re.sub(r"\s+", " ", text or "").strip()
    if len(compact) <= max_chars:
        return compact
    return compact[: max_chars - 1].rstrip() + "…"


def event_search_blob(event: Dict[str, Any]) -> str:
    parts = [
        event.get("event_id"),
        event.get("speaker"),
        event.get("speaker_label"),
        event.get("channel"),
        event.get("provider"),
        event.get("chat_id"),
        event.get("session_key"),
        event.get("clean_text"),
        event.get("raw_text"),
    ]
    for ptr in event.get("assets") or []:
        parts.extend([ptr.get("asset_id"), ptr.get("kind"), ptr.get("mime_type"), ptr.get("capture_status")])
    for ptr in event.get("deliveries") or []:
        parts.extend([ptr.get("delivery_id"), ptr.get("status"), ptr.get("target"), ptr.get("channel_message_id")])
    return " ".join(str(p) for p in parts if p is not None).lower()


def registry_search_blob(record: Dict[str, Any], keys: List[str]) -> str:
    values: List[str] = []
    for key in keys:
        value = record.get(key)
        if value is not None:
            values.append(str(value))
    values.append(json.dumps(record.get("source_ref") or {}, ensure_ascii=False, sort_keys=True))
    return " ".join(values).lower()


def evidence_search(ctx: RawOsContext, *, anchor_day: str, mainline: str, query: str, limit: int = 10) -> Dict[str, Any]:
    """Read-only keyword evidence search over ledger/assets/deliveries.

    This is deliberately not a default-memory path. It returns bounded snippets
    plus source refs so callers can explicitly drill down when they need proof.
    """
    q = (query or "").strip().lower()
    if not q:
        raise RawOsError("query must not be empty")
    events = sort_events_for_render(read_jsonl(ctx.ledger_path(anchor_day, mainline)))
    assets = load_asset_registry(ctx, anchor_day)
    deliveries = load_delivery_registry(ctx, anchor_day)
    matches: List[Dict[str, Any]] = []

    for event in events:
        if q not in event_search_blob(event):
            continue
        matches.append({
            "kind": "event",
            "event_id": event.get("event_id"),
            "timestamp": event.get("ts"),
            "speaker": event.get("speaker"),
            "speaker_label": event.get("speaker_label"),
            "snippet": text_snippet(event.get("clean_text") or event.get("raw_text") or ""),
            "source_ref": event.get("source_ref") or {},
            "asset_refs": event.get("assets") or [],
            "delivery_refs": event.get("deliveries") or [],
            "ranking_reason": "keyword match in ledger event",
        })

    for asset_id, record in assets.items():
        blob = registry_search_blob(record, ["asset_id", "source_event_id", "original_filename", "mime_type", "kind", "sha256", "storage_uri"])
        if q not in blob:
            continue
        matches.append({
            "kind": "asset",
            "asset_id": asset_id,
            "event_id": record.get("source_event_id"),
            "snippet": text_snippet(f"{record.get('original_filename') or asset_id} {record.get('mime_type') or ''} {(record.get('capture') or {}).get('status') or ''}"),
            "source_ref": record.get("source_ref") or {},
            "asset_refs": [{"asset_id": asset_id, "capture_status": (record.get("capture") or {}).get("status"), "storage_uri": record.get("storage_uri")}],
            "delivery_refs": [],
            "ranking_reason": "keyword match in asset registry",
        })

    for delivery_id, record in deliveries.items():
        blob = registry_search_blob(record, ["delivery_id", "source_event_id", "channel", "provider", "target", "chat_id", "channel_message_id", "status", "error"])
        if q not in blob:
            continue
        matches.append({
            "kind": "delivery",
            "delivery_id": delivery_id,
            "event_id": record.get("source_event_id"),
            "snippet": text_snippet(f"{record.get('status') or 'unknown'} {record.get('target') or ''} {record.get('channel_message_id') or ''} {record.get('error') or ''}"),
            "source_ref": record.get("source_ref") or {},
            "asset_refs": [],
            "delivery_refs": [{"delivery_id": delivery_id, "status": record.get("status"), "target": record.get("target"), "channel_message_id": record.get("channel_message_id")}],
            "ranking_reason": "keyword match in delivery registry",
        })

    ranked = []
    for rank, match in enumerate(matches[: max(0, limit)], start=1):
        ranked.append({"rank": rank, **match})
    return {
        "schema": SCHEMA_EVIDENCE_SEARCH_RESULT,
        "ok": True,
        "query": query,
        "anchor_day": anchor_day,
        "mainline": mainline,
        "limit": limit,
        "matches": ranked,
    }


def raw_replay_bundle(ctx: RawOsContext, *, anchor_day: str, mainline: str, event_ids: Optional[List[str]] = None, limit: Optional[int] = None) -> Dict[str, Any]:
    """Read-only replay bundle from ledger and registries, never rendered raw-md."""
    selected_ids = set(event_ids or [])
    events = sort_events_for_render(read_jsonl(ctx.ledger_path(anchor_day, mainline)))
    if selected_ids:
        events = [event for event in events if event.get("event_id") in selected_ids]
    if limit is not None:
        events = events[: max(0, limit)]
    assets_by_id = load_asset_registry(ctx, anchor_day)
    deliveries_by_id = load_delivery_registry(ctx, anchor_day)
    asset_ids = {ptr.get("asset_id") for event in events for ptr in event.get("assets") or [] if ptr.get("asset_id")}
    delivery_ids = {ptr.get("delivery_id") for event in events for ptr in event.get("deliveries") or [] if ptr.get("delivery_id")}
    issues: List[Dict[str, Any]] = []
    assets = []
    deliveries = []
    for asset_id in sorted(asset_ids):
        record = assets_by_id.get(asset_id)
        if record:
            assets.append(record)
        else:
            issues.append({"level": "error", "kind": "missing_asset_record", "asset_id": asset_id})
    for delivery_id in sorted(delivery_ids):
        record = deliveries_by_id.get(delivery_id)
        if record:
            deliveries.append(record)
        else:
            issues.append({"level": "error", "kind": "missing_delivery_record", "delivery_id": delivery_id})
    if selected_ids:
        found_ids = {event.get("event_id") for event in events}
        for missing_id in sorted(selected_ids - found_ids):
            issues.append({"level": "warning", "kind": "event_not_found", "event_id": missing_id})
    return {
        "schema": SCHEMA_REPLAY_BUNDLE,
        "ok": not any(issue.get("level") == "error" for issue in issues),
        "replay_id": stable_id("replay", anchor_day, mainline, ",".join(sorted(selected_ids)) if selected_ids else "all"),
        "window": {
            "anchor_day": anchor_day,
            "mainline": mainline,
            "event_ids": sorted(selected_ids),
            "limit": limit,
        },
        "events": events,
        "assets": assets,
        "deliveries": deliveries,
        "integrity": {
            "ledger_ok": True,
            "asset_refs_ok": not any(issue.get("kind") == "missing_asset_record" for issue in issues),
            "delivery_refs_ok": not any(issue.get("kind") == "missing_delivery_record" for issue in issues),
            "issues": issues,
        },
    }


def update_memory_collection_manifest(ctx: RawOsContext, *, anchor_day: str, mainline: str, memory_md: Path) -> Path:
    """Record raw-memory-md as an explicit evidence/memory candidate.

    OpenClaw memory ingestion is deployment-specific; this manifest is the
    stable Raw OS handoff: a collector can watch one file and import the listed
    projections without scraping directories heuristically.

    Raw-derived projections are not default conversational memory. Importers
    should keep them in an explicit evidence retrieval lane unless a deployment
    deliberately opts in to mixing them into default recall.
    """
    path = ctx.storage["stateDir"] / "memory-collection" / "raw-memory-md.json"
    current: Dict[str, Any] = {}
    if path.exists():
        try:
            current = read_json(path)
        except Exception:
            current = {}
    files = [f for f in current.get("files", []) if isinstance(f, dict)]
    rel = os.path.relpath(memory_md, ctx.root)
    files = [f for f in files if not (f.get("anchor_day") == anchor_day and f.get("mainline") == mainline)]
    files.append({"anchor_day": anchor_day, "mainline": mainline, "path": rel, "updated_at": utc_now_iso()})
    write_json(path, {
        "schema": "raw-memory-collection/v1",
        "kind": "raw-memory-md",
        "root": str(ctx.root),
        "default_recall": False,
        "retrieval_lane": "explicit-evidence",
        "importer_note": "Do not silently mix raw-derived projections into default conversational memory search; use explicit evidence retrieval unless opted in.",
        "files": sorted(files, key=lambda x: (x.get("anchor_day", ""), x.get("mainline", "")))
    })
    return path

def write_incident(ctx: RawOsContext, *, kind: str, error: str, context: Dict[str, Any]) -> Path:
    now = utc_now_iso()
    incident_id = stable_id("incident", kind, error, now)
    path = ctx.incident_dir() / f"{now.replace(':', '').replace('.', '_')}__{incident_id}.json"
    write_json(path, {
        "schema": "raw-incident/v1",
        "incident_id": incident_id,
        "kind": kind,
        "error": error,
        "context": context,
        "created_at": now,
    })
    return path


def audit_day(ctx: RawOsContext, *, anchor_day: str, mainline: str) -> Dict[str, Any]:
    ctx.assert_write_allowed()
    issues: List[Dict[str, Any]] = []
    ledger_path = ctx.ledger_path(anchor_day, mainline)
    registry_path = ctx.asset_registry_path(anchor_day)
    delivery_registry_path = ctx.delivery_registry_path(anchor_day)
    events = read_jsonl(ledger_path)
    registry = load_asset_registry(ctx, anchor_day)
    delivery_registry = load_delivery_registry(ctx, anchor_day)

    seen_event_ids: Dict[str, int] = {}
    referenced_assets = set()
    referenced_deliveries = set()
    for event in events:
        event_id = event.get("event_id")
        if event_id:
            seen_event_ids[event_id] = seen_event_ids.get(event_id, 0) + 1
        for ptr in event.get("deliveries") or []:
            did = ptr.get("delivery_id")
            if did:
                referenced_deliveries.add(did)
            record = delivery_registry.get(did)
            if not record:
                issues.append({"level": "error", "kind": "missing_delivery_record", "event_id": event.get("event_id"), "delivery_id": did})
                continue
            if not record.get("status"):
                issues.append({"level": "error", "kind": "delivery_status_missing", "delivery_id": did})
            if record.get("status") in {"failed", "cancelled"} and not record.get("error") and not record.get("fallback_reason"):
                issues.append({"level": "warning", "kind": "delivery_failure_without_reason", "delivery_id": did})
        for ptr in event.get("assets") or []:
            aid = ptr.get("asset_id")
            if aid:
                referenced_assets.add(aid)
            record = registry.get(aid)
            if not record:
                issues.append({"level": "error", "kind": "missing_asset_record", "event_id": event.get("event_id"), "asset_id": aid})
                continue
            capture = record.get("capture") or {}
            if capture.get("status") == "captured":
                uri = record.get("storage_uri")
                path = (ctx.root / uri).resolve() if uri else None
                if not path or not path.exists():
                    issues.append({"level": "error", "kind": "missing_asset_file", "asset_id": aid, "storage_uri": uri})
                elif record.get("sha256") and sha256_file(path) != record.get("sha256"):
                    issues.append({"level": "error", "kind": "asset_hash_mismatch", "asset_id": aid, "storage_uri": uri})
            elif not capture.get("error"):
                issues.append({"level": "warning", "kind": "capture_failure_without_reason", "asset_id": aid})

    for event_id, count in seen_event_ids.items():
        if count > 1:
            issues.append({"level": "warning", "kind": "duplicate_event_id", "event_id": event_id, "count": count})
    for asset_id in set(registry) - referenced_assets:
        issues.append({"level": "warning", "kind": "orphan_asset_record", "asset_id": asset_id})
    for delivery_id in set(delivery_registry) - referenced_deliveries:
        issues.append({"level": "warning", "kind": "orphan_delivery_record", "delivery_id": delivery_id})

    official = ctx.official_md_path(anchor_day, mainline)
    if not official.exists():
        issues.append({"level": "warning", "kind": "official_md_missing", "path": str(official)})

    result = {
        "schema": SCHEMA_AUDIT,
        "ok": not any(i.get("level") == "error" for i in issues),
        "anchor_day": anchor_day,
        "mainline": mainline,
        "ledger": str(ledger_path),
        "asset_registry": str(registry_path),
        "delivery_registry": str(delivery_registry_path),
        "events": len(events),
        "assets": len(registry),
        "deliveries": len(delivery_registry),
        "issues": issues,
        "audited_at": utc_now_iso(),
    }
    audit_path = ctx.storage["stateDir"] / "audits" / f"{anchor_day}__{mainline}.json"
    write_json(audit_path, result)
    result["audit_path"] = str(audit_path)
    return result


def cmd_validate(args: argparse.Namespace) -> int:
    config = read_json(Path(args.config))
    errors = config_errors(config)
    print(json.dumps({"ok": not errors, "errors": errors}, ensure_ascii=False, indent=2))
    return 1 if errors else 0


def cmd_init(args: argparse.Namespace) -> int:
    ctx = RawOsContext.load(Path(args.config))
    ctx.ensure_storage()
    print(json.dumps({"ok": True, "root": str(ctx.root), "storage": {k: str(v) for k, v in ctx.storage.items()}}, ensure_ascii=False, indent=2))
    return 0


def _doctor_add(checks: List[Dict[str, Any]], *, name: str, ok: bool, level: str, detail: str, remediation: Optional[str] = None) -> None:
    item: Dict[str, Any] = {"name": name, "ok": ok, "level": level, "detail": detail}
    if remediation:
        item["remediation"] = remediation
    checks.append(item)


def doctor_checks(ctx: RawOsContext, *, spool: Optional[str] = None) -> Dict[str, Any]:
    checks: List[Dict[str, Any]] = []

    py_ok = sys.version_info >= (3, 9)
    _doctor_add(
        checks,
        name="python_version",
        ok=py_ok,
        level="fatal",
        detail=f"{sys.version_info.major}.{sys.version_info.minor}.{sys.version_info.micro}",
        remediation="Install Python 3.9+ and run scripts/raw-os with that interpreter.",
    )

    for tool in ["bash", "date", "mkdir", "cp", "find"]:
        found = shutil.which(tool)
        _doctor_add(
            checks,
            name=f"shell_tool:{tool}",
            ok=bool(found),
            level="fatal",
            detail=found or "missing",
            remediation=f"Install {tool} or run on a POSIX-like Linux/macOS environment.",
        )

    tz_name = str((ctx.config.get("time") or {}).get("timezone") or "")
    try:
        ZoneInfo(tz_name)
        tz_ok = True
        tz_detail = tz_name
    except Exception as exc:
        tz_ok = False
        tz_detail = f"{tz_name or '<missing>'}: {exc}"
    _doctor_add(
        checks,
        name="timezone",
        ok=tz_ok,
        level="fatal",
        detail=tz_detail,
        remediation="Set time.timezone to a valid IANA timezone such as Asia/Shanghai.",
    )

    docx_available = importlib.util.find_spec("docx") is not None
    _doctor_add(
        checks,
        name="python_docx",
        ok=docx_available,
        level="warning",
        detail="available" if docx_available else "missing; built-in OOXML fallback will be used",
        remediation="Install python-docx for richer official docx rendering, or accept the built-in fallback.",
    )

    root_parent = ctx.root.parent
    _doctor_add(
        checks,
        name="workspace_parent",
        ok=root_parent.exists(),
        level="fatal",
        detail=str(root_parent),
        remediation="Create the parent directory or update agent.workspaceRoot.",
    )

    write_guard = None
    try:
        ctx.assert_write_allowed()
    except RawOsError as exc:
        write_guard = str(exc)
    _doctor_add(
        checks,
        name="write_guard",
        ok=write_guard is None,
        level="warning",
        detail=write_guard or "writes allowed",
        remediation="For an intentional live OpenClaw deploy, set RAW_OS_ALLOW_OPENCLAW_WORKSPACE=1. For protected production roots, also set RAW_OS_ALLOW_PROTECTED_WORKSPACE=1.",
    )

    for key, path in ctx.storage.items():
        parent = path.parent
        parent_ok = parent.exists() or ctx.root.exists() or root_parent.exists()
        _doctor_add(
            checks,
            name=f"storage_parent:{key}",
            ok=parent_ok,
            level="fatal",
            detail=str(parent),
            remediation="Run init after confirming the root, or create the parent directory.",
        )

    if spool:
        spool_path = Path(spool).expanduser().resolve()
        _doctor_add(
            checks,
            name="spool_readable",
            ok=spool_path.exists() and spool_path.is_file() and os.access(spool_path, os.R_OK),
            level="fatal",
            detail=str(spool_path),
            remediation="Create a normalized event JSONL spool or pass the correct --spool path.",
        )
        if spool_path.exists() and spool_path.is_file():
            invalid = None
            with spool_path.open("r", encoding="utf-8") as f:
                for line_no, line in enumerate(f, start=1):
                    line = line.strip().lstrip("\x00")
                    if not line:
                        continue
                    try:
                        obj = json.loads(line)
                    except json.JSONDecodeError as exc:
                        invalid = f"{spool_path}:{line_no}: {exc}"
                        break
                    if not isinstance(obj, dict):
                        invalid = f"{spool_path}:{line_no}: expected JSON object"
                        break
                    if not obj.get("event_id") or not (obj.get("timestamp") or obj.get("ts")):
                        invalid = f"{spool_path}:{line_no}: missing event_id or timestamp"
                        break
            _doctor_add(
                checks,
                name="spool_jsonl_shape",
                ok=invalid is None,
                level="fatal",
                detail=invalid or "valid JSONL event sample",
                remediation="Write one normalized event JSON object per line with event_id and timestamp.",
            )

    fatal = [c for c in checks if not c["ok"] and c["level"] == "fatal"]
    warnings = [c for c in checks if not c["ok"] and c["level"] == "warning"]
    return {
        "ok": not fatal,
        "checks": checks,
        "fatal_count": len(fatal),
        "warning_count": len(warnings),
        "write_guard": write_guard,
        "is_openclaw_workspace": ctx.is_openclaw_workspace(),
        "is_protected_workspace": ctx.root in protected_live_workspace_roots(),
    }


def cmd_doctor(args: argparse.Namespace) -> int:
    ctx = RawOsContext.load(Path(args.config))
    doctor = doctor_checks(ctx, spool=getattr(args, "spool", None))
    result = {
        "ok": doctor["ok"],
        "config": str(ctx.config_path),
        "root": str(ctx.root),
        "agent_id": ctx.agent_id,
        "is_openclaw_workspace": doctor["is_openclaw_workspace"],
        "is_protected_workspace": doctor["is_protected_workspace"],
        "write_allowed": doctor["write_guard"] is None,
        "write_guard": doctor["write_guard"],
        "fatal_count": doctor["fatal_count"],
        "warning_count": doctor["warning_count"],
        "checks": doctor["checks"],
        "storage": {k: str(v) for k, v in ctx.storage.items()},
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["ok"] else 1


def cmd_ingest(args: argparse.Namespace) -> int:
    ctx = RawOsContext.load(Path(args.config))
    result = ingest_event(
        ctx,
        anchor_day=args.anchor_day,
        mainline=args.mainline,
        text=args.text or "",
        speaker=args.speaker,
        sender_id=args.sender_id,
        sender_label=args.sender_label,
        channel=args.channel,
        provider=args.provider or args.channel,
        chat_id=args.chat_id,
        message_id=args.message_id,
        session_key=args.session_key,
        attachments=[Path(p) for p in args.attachment or []],
        failed_assets=args.failed_asset or [],
        failure_reason=args.failure_reason,
        metadata=parse_kv_items(args.meta),
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def cmd_ingest_normalized(args: argparse.Namespace) -> int:
    ctx = RawOsContext.load(Path(args.config))
    event = read_json(Path(args.event))
    result = ingest_normalized_event(ctx, anchor_day=args.anchor_day, mainline=args.mainline, event=event)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def ingest_tap_spool_file(ctx: RawOsContext, args: argparse.Namespace) -> Dict[str, Any]:
    spool = Path(args.spool).expanduser().resolve()
    if not spool.exists() or not spool.is_file():
        raise RawOsError(f"tap spool not found: {spool}")
    state_path = ctx.tap_spool_state_path(args.anchor_day, args.mainline)
    state: Dict[str, Any] = {}
    processed_ids = set()
    if args.stateful and state_path.exists():
        state = read_json(state_path)
        processed_ids = set(state.get("processed_event_ids") or [])

    ingested = 0
    skipped = 0
    duplicates = 0
    state_skipped = 0
    event_ids: List[str] = []
    newly_processed: List[str] = []
    with spool.open("r", encoding="utf-8") as f:
        for line_no, line in enumerate(f, start=1):
            # Some append paths can leave NUL padding before an otherwise valid
            # JSON object; salvage the record instead of blocking the whole day.
            line = line.strip().lstrip("\x00")
            if not line:
                continue
            try:
                event = json.loads(line)
            except json.JSONDecodeError as exc:
                raise RawOsError(f"invalid tap spool JSON at {spool}:{line_no}: {exc}") from exc
            event_id = str(event.get("event_id") or f"{spool}:{line_no}")
            if args.stateful and event_id in processed_ids:
                state_skipped += 1
                skipped += 1
                continue
            if not event_matches_anchor_day(event.get("timestamp") or event.get("ts"), args.anchor_day, ctx.config):
                skipped += 1
                continue
            result = ingest_normalized_event(ctx, anchor_day=args.anchor_day, mainline=args.mainline, event=event)
            if result.get("skipped"):
                duplicates += 1
            else:
                ingested += 1
            event_ids.append(str(result.get("event_id") or event_id))
            newly_processed.append(event_id)

    if args.stateful:
        ensure_dir(state_path.parent)
        all_processed = sorted(processed_ids.union(newly_processed))
        write_json(state_path, {
            "schema": "raw-os-tap-spool-state/v1",
            "anchor_day": args.anchor_day,
            "mainline": args.mainline,
            "spool": str(spool),
            "updated_at": utc_now_iso(),
            "processed_event_ids": all_processed,
        })
    return {
        "ok": True,
        "spool": str(spool),
        "ingested": ingested,
        "skipped": skipped,
        "duplicates": duplicates,
        "state_skipped": state_skipped,
        "state_path": str(state_path),
        "event_ids": event_ids,
    }


def cmd_ingest_tap_spool(args: argparse.Namespace) -> int:
    ctx = RawOsContext.load(Path(args.config))
    result = ingest_tap_spool_file(ctx, args)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def cmd_daily(args: argparse.Namespace) -> int:
    ctx = RawOsContext.load(Path(args.config))
    result: Dict[str, Any] = {
        "ok": True,
        "anchor_day": args.anchor_day,
        "mainline": args.mainline,
        "send": bool(args.send),
    }
    if args.send:
        raise RawOsError("daily --send is not implemented yet; run delivery through a deployment adapter")
    if args.spool:
        ingest_args = argparse.Namespace(
            anchor_day=args.anchor_day,
            mainline=args.mainline,
            spool=args.spool,
            stateful=args.stateful,
        )
        result["ingest"] = ingest_tap_spool_file(ctx, ingest_args)
    result["render"] = render_day(ctx, anchor_day=args.anchor_day, mainline=args.mainline)
    if not args.no_docx:
        result["docx"] = render_official_docx(ctx, anchor_day=args.anchor_day, mainline=args.mainline)
    result["memory"] = render_memory_day(ctx, anchor_day=args.anchor_day, mainline=args.mainline)
    audit = audit_day(ctx, anchor_day=args.anchor_day, mainline=args.mainline)
    result["audit"] = audit
    result["ok"] = bool(audit.get("ok"))
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["ok"] else 1


def ingest_transcript_file(ctx: RawOsContext, args: argparse.Namespace) -> Dict[str, Any]:
    transcript = Path(args.transcript).expanduser().resolve()
    if not transcript.exists() or not transcript.is_file():
        raise RawOsError(f"transcript not found: {transcript}")
    state_path = ctx.transcript_state_path(args.anchor_day, args.mainline)
    state: Dict[str, Any] = {}
    processed_records = set()
    if args.stateful and state_path.exists():
        state = read_json(state_path)
        processed_by_transcript = state.get("processed_by_transcript") or {}
        if processed_by_transcript:
            processed_records = set(processed_by_transcript.get(str(transcript)) or [])
        else:
            processed_records = set(state.get("processed_record_ids") or [])

    count = 0
    skipped = 0
    duplicates = 0
    state_skipped = 0
    event_ids: List[str] = []
    newly_processed: List[str] = []
    transcript_records: List[Dict[str, Any]] = []
    runtime_context_by_parent: Dict[str, str] = {}
    with transcript.open("r", encoding="utf-8") as f:
        for line_no, line in enumerate(f, start=1):
            line = line.strip()
            if not line:
                continue
            try:
                obj = json.loads(line)
            except json.JSONDecodeError as exc:
                raise RawOsError(f"invalid transcript json {transcript}:{line_no}: {exc}") from exc
            obj["_raw_os_line_no"] = line_no
            transcript_records.append(obj)
            if obj.get("type") == "custom_message" and obj.get("customType") == "openclaw.runtime-context" and obj.get("parentId"):
                runtime_context_by_parent[str(obj.get("parentId"))] = str(obj.get("content") or "")

    for obj in transcript_records:
            line_no = int(obj.get("_raw_os_line_no") or 0)
            record_id = str(obj.get("id") or f"line:{line_no}")
            if args.stateful and record_id in processed_records:
                state_skipped += 1
                skipped += 1
                continue
            obj["session_file"] = str(transcript)
            event = normalized_event_from_transcript_message(
                obj,
                default_channel=args.default_channel,
                default_provider=args.default_provider,
                default_sender_id=args.sender_id,
                default_chat_id=args.chat_id,
                runtime_context=runtime_context_by_parent.get(record_id),
            )
            if not event:
                skipped += 1
                continue
            if not event_matches_anchor_day(event.get("timestamp") or event.get("ts"), args.anchor_day, ctx.config):
                skipped += 1
                continue
            if args.sender_id and event.get("speaker") == "user" and sender_id_from_event(event) != args.sender_id:
                skipped += 1
                continue
            if args.chat_id and event.get("chat_id") != args.chat_id:
                skipped += 1
                continue
            result = ingest_normalized_event(ctx, anchor_day=args.anchor_day, mainline=args.mainline, event=event)
            newly_processed.append(record_id)
            if result.get("skipped"):
                duplicates += 1
                skipped += 1
                continue
            count += 1
            event_ids.append(result["event_id"])

    if args.stateful:
        merged = sorted(processed_records.union(newly_processed))
        processed_by_transcript = dict(state.get("processed_by_transcript") or {})
        processed_by_transcript[str(transcript)] = merged
        all_processed = sorted({rid for ids in processed_by_transcript.values() if isinstance(ids, list) for rid in ids})
        state = {
            "schema": "raw-transcript-ingest-state/v1",
            "anchor_day": args.anchor_day,
            "mainline": args.mainline,
            "transcript": str(transcript),
            "processed_by_transcript": processed_by_transcript,
            "sender_id": args.sender_id,
            "chat_id": args.chat_id,
            "processed_record_ids": all_processed,
            "updated_at": utc_now_iso(),
            "last_result": {
                "ingested": count,
                "skipped": skipped,
                "duplicates": duplicates,
                "state_skipped": state_skipped,
                "event_ids": event_ids,
            },
        }
        write_json(state_path, state)

    return {"ok": True, "transcript": str(transcript), "ingested": count, "skipped": skipped, "duplicates": duplicates, "state_skipped": state_skipped, "state_path": str(state_path) if args.stateful else None, "event_ids": event_ids}


def cmd_ingest_transcript(args: argparse.Namespace) -> int:
    ctx = RawOsContext.load(Path(args.config))
    result = ingest_transcript_file(ctx, args)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def run_tick(
    ctx: RawOsContext,
    *,
    anchor_day: str,
    mainline: str,
    transcript: Path,
    sender_id: Optional[str],
    chat_id: Optional[str],
    default_channel: str,
    default_provider: str,
    dry_run: bool = False,
) -> Dict[str, Any]:
    if dry_run:
        candidates = 0
        skipped = 0
        records: List[Dict[str, Any]] = []
        runtime_context_by_parent: Dict[str, str] = {}
        with transcript.open("r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                try:
                    obj = json.loads(line)
                except json.JSONDecodeError:
                    skipped += 1
                    continue
                records.append(obj)
                if obj.get("type") == "custom_message" and obj.get("customType") == "openclaw.runtime-context" and obj.get("parentId"):
                    runtime_context_by_parent[str(obj.get("parentId"))] = str(obj.get("content") or "")
        for obj in records:
                event = normalized_event_from_transcript_message(
                    obj,
                    default_channel=default_channel,
                    default_provider=default_provider,
                    default_sender_id=sender_id,
                    default_chat_id=chat_id,
                    runtime_context=runtime_context_by_parent.get(str(obj.get("id"))),
                )
                if not event:
                    skipped += 1
                    continue
                if not event_matches_anchor_day(event.get("timestamp") or event.get("ts"), anchor_day, ctx.config):
                    skipped += 1
                    continue
                if sender_id and event.get("speaker") == "user" and sender_id_from_event(event) != sender_id:
                    skipped += 1
                    continue
                if chat_id and event.get("chat_id") != chat_id:
                    skipped += 1
                    continue
                candidates += 1
        return {"ok": True, "dry_run": True, "transcript": str(transcript), "candidates": candidates, "skipped": skipped}

    ingest_args = argparse.Namespace(
        config=str(ctx.config_path),
        anchor_day=anchor_day,
        mainline=mainline,
        transcript=str(transcript),
        sender_id=sender_id,
        chat_id=chat_id,
        default_channel=default_channel,
        default_provider=default_provider,
        stateful=True,
    )
    # Reuse the command implementation but return structured phase results by
    # calling the underlying functions after ingest completes.
    ingest = ingest_transcript_file(ctx, ingest_args)
    render = render_day(ctx, anchor_day=anchor_day, mainline=mainline)
    docx = None
    delivery_config = ctx.config.get("delivery") or {}
    if str(delivery_config.get("officialFormat") or "md").lower() == "docx" or delivery_config.get("alsoGenerateOfficialDocx"):
        docx = render_official_docx(ctx, anchor_day=anchor_day, mainline=mainline)
    memory = render_memory_day(ctx, anchor_day=anchor_day, mainline=mainline)
    audit = audit_day(ctx, anchor_day=anchor_day, mainline=mainline)
    incident_path = None
    if not audit.get("ok"):
        incident_path = write_incident(ctx, kind="audit_error", error="audit failed after tick", context={"anchor_day": anchor_day, "mainline": mainline, "transcript": str(transcript), "audit": audit})
    return {"ok": audit.get("ok", False), "dry_run": False, "transcript": str(transcript), "ingest": ingest, "render": render, "docx": docx, "memory": memory, "audit": audit, "incident_path": str(incident_path) if incident_path else None}


def cmd_tick(args: argparse.Namespace) -> int:
    ctx = RawOsContext.load(Path(args.config))
    transcript = Path(args.transcript).expanduser().resolve()
    try:
        if not transcript.exists() or not transcript.is_file():
            raise RawOsError(f"transcript not found: {transcript}")
        result = run_tick(
            ctx,
            anchor_day=args.anchor_day,
            mainline=args.mainline,
            transcript=transcript,
            sender_id=args.sender_id,
            chat_id=args.chat_id,
            default_channel=args.default_channel,
            default_provider=args.default_provider,
            dry_run=args.dry_run,
        )
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0 if result.get("ok") else 1
    except Exception as exc:
        incident_path = None
        try:
            incident_path = write_incident(ctx, kind="tick_error", error=str(exc), context={
                "anchor_day": args.anchor_day,
                "mainline": args.mainline,
                "transcript": str(transcript),
                "dry_run": args.dry_run,
            })
        except Exception:
            pass
        print(json.dumps({"ok": False, "error": str(exc), "incident_path": str(incident_path) if incident_path else None}, ensure_ascii=False, indent=2), file=sys.stderr)
        return 1


def cmd_render(args: argparse.Namespace) -> int:
    ctx = RawOsContext.load(Path(args.config))
    print(json.dumps(render_day(ctx, anchor_day=args.anchor_day, mainline=args.mainline), ensure_ascii=False, indent=2))
    return 0


def cmd_render_official_docx(args: argparse.Namespace) -> int:
    ctx = RawOsContext.load(Path(args.config))
    print(json.dumps(render_official_docx(ctx, anchor_day=args.anchor_day, mainline=args.mainline), ensure_ascii=False, indent=2))
    return 0


def cmd_audit(args: argparse.Namespace) -> int:
    ctx = RawOsContext.load(Path(args.config))
    result = audit_day(ctx, anchor_day=args.anchor_day, mainline=args.mainline)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["ok"] else 1



def cmd_render_memory(args: argparse.Namespace) -> int:
    ctx = RawOsContext.load(Path(args.config))
    print(json.dumps(render_memory_day(ctx, anchor_day=args.anchor_day, mainline=args.mainline), ensure_ascii=False, indent=2))
    return 0


def cmd_evidence_search(args: argparse.Namespace) -> int:
    ctx = RawOsContext.load(Path(args.config))
    result = evidence_search(ctx, anchor_day=args.anchor_day, mainline=args.mainline, query=args.query, limit=args.limit)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def cmd_raw_replay(args: argparse.Namespace) -> int:
    ctx = RawOsContext.load(Path(args.config))
    result = raw_replay_bundle(ctx, anchor_day=args.anchor_day, mainline=args.mainline, event_ids=args.event_id, limit=args.limit)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result.get("ok") else 1


def cmd_asset_path(args: argparse.Namespace) -> int:
    ctx = RawOsContext.load(Path(args.config))
    record, registry_path = find_asset_record(ctx, asset_id=args.asset_id, anchor_day=args.anchor_day)
    path = resolve_asset_path(ctx, record)
    print(json.dumps({"ok": True, "asset_id": args.asset_id, "path": str(path), "registry": str(registry_path), "record": record}, ensure_ascii=False, indent=2))
    return 0


def cmd_asset_export(args: argparse.Namespace) -> int:
    ctx = RawOsContext.load(Path(args.config))
    print(json.dumps(export_asset(ctx, asset_id=args.asset_id, out_dir=Path(args.out).expanduser(), anchor_day=args.anchor_day), ensure_ascii=False, indent=2))
    return 0


def cmd_asset_resolve_request(args: argparse.Namespace) -> int:
    ctx = RawOsContext.load(Path(args.config))
    asset_id = args.asset_id or asset_id_from_natural_language(args.text or "")
    if not asset_id:
        print(json.dumps({"ok": False, "reason": "no_asset_id_found"}, ensure_ascii=False, indent=2))
        return 1
    record, registry_path = find_asset_record(ctx, asset_id=asset_id, anchor_day=args.anchor_day)
    result: Dict[str, Any] = {"ok": True, "asset_id": asset_id, "registry": str(registry_path), "record": record}
    if (record.get("capture") or {}).get("status") == "captured":
        result["path"] = str(resolve_asset_path(ctx, record))
    else:
        result["ok"] = False
        result["reason"] = "asset_not_captured"
        result["capture_error"] = (record.get("capture") or {}).get("error")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result.get("ok") else 1

def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="Raw OS v0 CLI")
    sub = p.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("validate-config")
    s.add_argument("--config", required=True)
    s.set_defaults(func=cmd_validate)

    s = sub.add_parser("init")
    s.add_argument("--config", required=True)
    s.set_defaults(func=cmd_init)

    s = sub.add_parser("doctor")
    s.add_argument("--config", required=True)
    s.add_argument("--spool", help="optional normalized event JSONL spool to validate before daily ingest")
    s.set_defaults(func=cmd_doctor)

    s = sub.add_parser("ingest-event")
    s.add_argument("--config", required=True)
    s.add_argument("--anchor-day", required=True)
    s.add_argument("--mainline", required=True)
    s.add_argument("--text")
    s.add_argument("--speaker", default="user")
    s.add_argument("--sender-id", required=True)
    s.add_argument("--sender-label")
    s.add_argument("--channel", required=True)
    s.add_argument("--provider")
    s.add_argument("--chat-id", required=True)
    s.add_argument("--message-id", required=True)
    s.add_argument("--session-key", default="raw-os:manual")
    s.add_argument("--attachment", action="append", help="local attachment path to capture")
    s.add_argument("--failed-asset", action="append", help="asset ref that failed capture")
    s.add_argument("--failure-reason")
    s.add_argument("--meta", action="append", help="extra metadata key=value")
    s.set_defaults(func=cmd_ingest)

    s = sub.add_parser("ingest-normalized")
    s.add_argument("--config", required=True)
    s.add_argument("--anchor-day", required=True)
    s.add_argument("--mainline", required=True)
    s.add_argument("--event", required=True, help="normalized event JSON file")
    s.set_defaults(func=cmd_ingest_normalized)

    s = sub.add_parser("ingest-tap-spool")
    s.add_argument("--config", required=True)
    s.add_argument("--anchor-day", required=True)
    s.add_argument("--mainline", required=True)
    s.add_argument("--spool", required=True, help="Raw OS tap spool JSONL written by openclaw-plugin-raw-os-tap")
    s.add_argument("--stateful", action="store_true", help="persist processed tap event ids under stateDir")
    s.set_defaults(func=cmd_ingest_tap_spool)

    s = sub.add_parser("ingest-spool", help="harness-neutral alias for ingest-tap-spool")
    s.add_argument("--config", required=True)
    s.add_argument("--anchor-day", required=True)
    s.add_argument("--mainline", required=True)
    s.add_argument("--spool", required=True, help="normalized event JSONL spool written by a harness adapter")
    s.add_argument("--stateful", action="store_true", help="persist processed event ids under stateDir")
    s.set_defaults(func=cmd_ingest_tap_spool)

    s = sub.add_parser("daily")
    s.add_argument("--config", required=True)
    s.add_argument("--anchor-day", required=True)
    s.add_argument("--mainline", required=True)
    s.add_argument("--spool", help="optional normalized event JSONL spool to ingest before rendering")
    s.add_argument("--stateful", action="store_true", help="persist processed spool event ids under stateDir")
    s.add_argument("--no-docx", action="store_true", help="skip official docx generation")
    s.add_argument("--send", action="store_true", help="reserved for deployment delivery adapters; currently unsupported")
    s.set_defaults(func=cmd_daily)

    s = sub.add_parser("ingest-transcript")
    s.add_argument("--config", required=True)
    s.add_argument("--anchor-day", required=True)
    s.add_argument("--mainline", required=True)
    s.add_argument("--transcript", required=True, help="OpenClaw session transcript .jsonl")
    s.add_argument("--sender-id", help="optional sender id filter")
    s.add_argument("--chat-id", help="optional chat id filter")
    s.add_argument("--default-channel", default="telegram")
    s.add_argument("--default-provider", default="telegram")
    s.add_argument("--stateful", action="store_true", help="persist processed transcript record ids under stateDir")
    s.set_defaults(func=cmd_ingest_transcript)

    s = sub.add_parser("tick")
    s.add_argument("--config", required=True)
    s.add_argument("--anchor-day", required=True)
    s.add_argument("--mainline", required=True)
    s.add_argument("--transcript", required=True)
    s.add_argument("--sender-id")
    s.add_argument("--chat-id")
    s.add_argument("--default-channel", default="telegram")
    s.add_argument("--default-provider", default="telegram")
    s.add_argument("--dry-run", action="store_true")
    s.set_defaults(func=cmd_tick)

    s = sub.add_parser("render-day")
    s.add_argument("--config", required=True)
    s.add_argument("--anchor-day", required=True)
    s.add_argument("--mainline", required=True)
    s.set_defaults(func=cmd_render)

    s = sub.add_parser("render-official-docx")
    s.add_argument("--config", required=True)
    s.add_argument("--anchor-day", required=True)
    s.add_argument("--mainline", required=True)
    s.set_defaults(func=cmd_render_official_docx)

    s = sub.add_parser("audit-day")
    s.add_argument("--config", required=True)
    s.add_argument("--anchor-day", required=True)
    s.add_argument("--mainline", required=True)
    s.set_defaults(func=cmd_audit)

    s = sub.add_parser("render-memory-day")
    s.add_argument("--config", required=True)
    s.add_argument("--anchor-day", required=True)
    s.add_argument("--mainline", required=True)
    s.set_defaults(func=cmd_render_memory)

    s = sub.add_parser("evidence-search")
    s.add_argument("--config", required=True)
    s.add_argument("--anchor-day", required=True)
    s.add_argument("--mainline", required=True)
    s.add_argument("--query", required=True)
    s.add_argument("--limit", type=int, default=10)
    s.set_defaults(func=cmd_evidence_search)

    s = sub.add_parser("raw-replay")
    s.add_argument("--config", required=True)
    s.add_argument("--anchor-day", required=True)
    s.add_argument("--mainline", required=True)
    s.add_argument("--event-id", action="append", help="event id to include; may be repeated")
    s.add_argument("--limit", type=int, help="maximum events when no event id is provided")
    s.set_defaults(func=cmd_raw_replay)

    s = sub.add_parser("asset-path")
    s.add_argument("--config", required=True)
    s.add_argument("--asset-id", required=True)
    s.add_argument("--anchor-day")
    s.set_defaults(func=cmd_asset_path)

    s = sub.add_parser("asset-export")
    s.add_argument("--config", required=True)
    s.add_argument("--asset-id", required=True)
    s.add_argument("--out", required=True)
    s.add_argument("--anchor-day")
    s.set_defaults(func=cmd_asset_export)

    s = sub.add_parser("asset-resolve-request")
    s.add_argument("--config", required=True)
    s.add_argument("--text", help="natural-language request containing an asset_id")
    s.add_argument("--asset-id", help="explicit asset id override")
    s.add_argument("--anchor-day")
    s.set_defaults(func=cmd_asset_resolve_request)
    return p


def main(argv: Optional[List[str]] = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        return int(args.func(args))
    except RawOsError as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False, indent=2), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
